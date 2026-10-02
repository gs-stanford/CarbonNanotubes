#!/usr/bin/env python3
"""Import an adjudicated review queue without promoting unchecked source rows.

The returned workbook is read-only. The generated JSON is the durable input to
the release builder; neither the workbook nor review PDFs are needed to rebuild.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
CANDIDATES = ROOT / "data/literature/review_extraction/normalized/review_measurement_candidates.csv"

# Factors convert source units into the release's SI fields, not display units.
PROPERTIES = {
    "tensile_strength": ("tensile_strength_Pa", {"MPa": 1e6, "GPa": 1e9}),
    "youngs_modulus": ("initial_modulus_Pa", {"GPa": 1e9}),
    "electrical_conductivity": ("electrical_conductivity_S_m", {"S cm^-1": 100, "S m^-1": 1, "10^6 S m^-1": 1e6}),
    "thermal_conductivity": ("thermal_conductivity_W_mK", {"W m^-1 K^-1": 1}),
    "density": ("density_kg_m3", {"g cm^-3": 1000}),
    "diameter": ("diameter_m", {"um": 1e-6}),
    "linear_density": ("linear_density_kg_m", {"tex": 1e-6}),
    "specific_strength": ("specific_strength_N_m_kg", {"GPa SG^-1": 1e6}),
    "specific_modulus": ("specific_modulus_N_m_kg", {"GPa SG^-1": 1e6}),
    "specific_electrical_conductivity": ("specific_electrical_conductivity_S_m2_kg", {"(S m^-1)/(g cm^-3)": 1e-3}),
    "ampacity": ("ampacity_A_m2", {"A cm^-2": 1e4}),
    "toughness": ("rupture_work_J_kg", {"J g^-1": 1000}),
    "electrical_resistivity": ("electrical_conductivity_S_m", {"mOhm cm": 1e-5}),
}
FAMILIES = {"CNT": "CNT_or_CNT_hybrid", "Graphene": "graphene_or_GO_fiber",
            "Carbon fiber": "carbon_fiber_comparator", "GIC": "other_carbon_comparator"}
ALIASES = {"Young's modulus": "youngs_modulus"}


def text(value: Any) -> str:
    return "" if value is None or pd.isna(value) else str(value).strip()


def json_value(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    return None if value is None or pd.isna(value) else value


def unit_key(value: Any) -> str:
    return text(value).translate(str.maketrans({"⁻": "^-", "¹": "1", "²": "2", "³": "3", "µ": "u", "μ": "u", "Ω": "Ohm"})).replace("S/cm", "S cm^-1")


def property_key(value: Any) -> str:
    return ALIASES.get(text(value), text(value).lower().replace(" ", "_"))


def number(value: Any) -> float | None:
    try:
        result = float(value)
        return result if math.isfinite(result) else None
    except (ValueError, TypeError):
        return None


def parse_expression(raw: str) -> tuple[float, float | None, float | None, float | None]:
    """Parse positive source scalars, bounds, ranges and reported +/- errors."""
    value = re.sub(r"\s+", "", raw).replace("−", "-").replace("–", "-").replace("×", "x")
    value = re.sub(r"^(approximately|approx\.?|[<>~])", "", value, flags=re.I)
    value = re.sub(r"a\)?$", "", value)
    exponent = re.search(r"(?:x10|^10)(?:\^\(?([+-]?\d+)\)?|([+-]?\d+))$", value)
    scale = 1.0
    if exponent:
        scale = 10 ** int(exponent.group(1) or exponent.group(2))
        value = value[:exponent.start()] or "1"
    scalar = r"(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?"
    match = re.fullmatch(rf"({scalar})(?:[-~]({scalar}))?(?:±({scalar}))?", value)
    if not match:
        raise ValueError(f"Unquantified or unsupported source expression: {raw}")
    first = float(match.group(1)) * scale
    last = float(match.group(2)) * scale if match.group(2) else None
    error = float(match.group(3)) * scale if match.group(3) else None
    if last is not None:
        if last < first:
            raise ValueError("Reversed source range")
        return (first + last) / 2, first, last, error
    return first, None, None, error


def normalized_measurement(row: dict, original: dict) -> tuple[float, dict]:
    prop = property_key(row["Property"])
    field, factors = PROPERTIES[prop]
    raw, unit = text(row["Reported value"]), unit_key(row["Reported unit"])
    if unit not in factors:
        raise ValueError(f"Unsupported {prop} source unit: {unit}")
    factor = factors[unit]
    edited = raw != text(original["raw_value"])
    direct, minimum, maximum, uncertainty = parse_expression(raw)
    value = direct * factor
    minimum = minimum * factor if minimum is not None else None
    maximum = maximum * factor if maximum is not None else None
    uncertainty = uncertainty * factor if uncertainty is not None else None
    if value <= 0 or not math.isfinite(value):
        raise ValueError("Canonical value must be finite and positive")
    if prop == "electrical_resistivity":
        if minimum is not None or uncertainty is not None:
            raise ValueError("Non-scalar resistivity must be reviewed separately")
        value = 1 / value
    range_midpoint = minimum is not None
    if range_midpoint and uncertainty is None:
        uncertainty = (maximum - minimum) / 2
    qualifier = raw.lstrip()
    bound = "range_midpoint" if range_midpoint else "point_estimate"
    if qualifier.startswith(">"):
        bound = "lower_bound"
    elif qualifier.startswith("<"):
        bound = "upper_bound"
    normalized = prop.startswith("specific_")
    derived = range_midpoint or prop == "electrical_resistivity"
    metadata = {
        "reported_value": None if derived else direct,
        "reported_expression": raw, "reported_unit": text(row["Reported unit"]),
        "reported_or_derived": "derived" if derived else "reported",
        "uncertainty_value_reported": uncertainty / factor if uncertainty is not None else None,
        "uncertainty_value_canonical": uncertainty,
        "uncertainty_type": "range" if range_midpoint else ("reported_unspecified" if uncertainty is not None else "not_reported"),
        "canonical_min": minimum, "canonical_max": maximum,
        "normalization_basis": "directly_reported_mass_specific" if normalized else "not_applicable",
        "density_basis": "unknown" if normalized else "not_applicable",
        "density_value_kg_m3": None, "density_source_locator": None,
        "cross_section_method": None, "value_bound_type": bound,
        "statistic_type": "unspecified", "source_locator": text(row["Source table / location"]),
        "derivation_formula": "value = (canonical_min + canonical_max) / 2" if range_midpoint else ("electrical_conductivity_S_m = 1 / electrical_resistivity_Ohm_m" if prop == "electrical_resistivity" else None),
        "derivation_inputs_json": json.dumps({"canonical_min": minimum, "canonical_max": maximum}) if range_midpoint else (json.dumps({"electrical_resistivity_Ohm_m": direct * factor}) if prop == "electrical_resistivity" else None),
        "review_candidate_id": text(row["Record ID"]),
        "source_conversion_formula": text(original.get("conversion_formula")),
        "reviewer": text(row["Reviewer"]), "review_date": json_value(row["Review date"]),
        "reviewer_notes": text(row["Reviewer notes"]), "raw_value_corrected": edited,
    }
    return value, metadata


def import_workbook(workbook: Path, candidate_path: Path = CANDIDATES) -> dict:
    candidates = pd.read_csv(candidate_path).set_index("candidate_id").to_dict("index")
    sheet = load_workbook(workbook, data_only=True)["Validation Queue"]
    headers = [c.value for c in sheet[8]]
    groups: dict[str, dict] = {}
    decisions, corrections, counts = [], [], Counter()
    for r in range(9, sheet.max_row + 1):
        row = dict(zip(headers, [sheet.cell(r, c).value for c in range(1, len(headers) + 1)]))
        if text(row.get("Row source")) != "New review candidate":
            # The reviewer accidentally filled some invalid rows across all
            # columns. Count their decision, but never treat those as samples.
            if r <= 402 and text(row.get("Human validation")):
                counts[text(row["Human validation"])] += 1
            continue
        status = text(row["Human validation"])
        counts[status] += 1
        candidate_id = text(row["Record ID"])
        decision = {"candidate_id": candidate_id, "workbook_row": r, "status": status,
                    "primary_source_opened": text(row["Primary source opened"]),
                    "duplicate_decision": text(row["Duplicate decision"]),
                    "reviewer_notes": text(row["Reviewer notes"])}
        decisions.append(decision)
        if status != "✓ Valid":
            decision["import_decision"] = "not_approved"
            continue
        if text(row["Primary source opened"]) != "✓ Opened" or text(row["Duplicate decision"]) not in {"No duplicate", "Different specimen"}:
            decision["import_decision"] = "held_incomplete_source_or_duplicate_review"
            continue
        original = candidates.get(candidate_id)
        if not original:
            raise ValueError(f"Approved row {r} has no original extraction identity")
        prop = property_key(row["Property"])
        doi = ";".join(part.strip() for part in text(row["Primary DOI"]).lower().split(";"))
        if not re.fullmatch(r"10\.\d{4,9}/[^\s;]+(?:;10\.\d{4,9}/[^\s;]+)*", doi):
            decision["import_decision"] = "held_missing_or_malformed_doi"
            continue
        family = FAMILIES.get(text(row["Material family"]))
        if family is None:
            raise ValueError(f"Unsupported family in approved row {r}")
        group_id = candidate_id.removesuffix("-" + original["property_name"])
        # Changes to DOI identity cannot silently merge a different publication.
        group_key = group_id + "|" + doi
        record = groups.setdefault(group_key, {
            "record_id": "rec_" + hashlib.sha1(("human-validation-v1|" + group_key).encode()).hexdigest()[:12],
            "source_file": "cnt_atlas_human_validation.xlsx", "source_sheet": "Validation Queue", "source_row": r,
            "record_label": text(row["Sample / material"]), "sample_name": text(row["Sample / material"]),
            "material_family": family, "material_subclass": text(row["Material subclass"]),
            "form_factor": "bulk" if text(row["Material family"]) == "GIC" else "fiber_yarn",
            "citation_raw": "; ".join("https://doi.org/" + d for d in doi.split(";")),
            "doi_raw": doi,
            "publication_title": text(row["Publication / review title"]),
            "compilation_source_doi_raw": text(original["review_doi"]),
            "compilation_source_title": text(original["review_title"]),
            "dataset_provenance": "human_validated_review_table_extraction",
            "primary_source_verification_status": "verified_against_primary_source" if ";" not in doi else "reviewed_multi_source_table_row",
            "source_citation_class": "peer_reviewed_research_record" if family in {"CNT_or_CNT_hybrid", "graphene_or_GO_fiber"} else "peer_reviewed_comparator",
            "comparison_scope": "strict_candidate_missing_conditions" if family in {"CNT_or_CNT_hybrid", "graphene_or_GO_fiber"} else "contextual_comparator",
            "specimen_linkage": "single_source_row_unverified" if ";" not in doi else "aggregated_across_specimens",
            "provenance_table_figure_page": text(row["Source table / location"]),
            "extraction_method": "human_validated_review_table", "curator": text(row["Reviewer"]),
            "human_duplicate_decision": text(row["Duplicate decision"]),
            "source_priority": 0, "values": {}, "measurement_metadata": {}, "validation_rows": [],
            "context_measurements": [], "conditions_notes": text(row["Conditions / notes"]),
        })
        record["validation_rows"].append({k: json_value(v) for k, v in row.items()})
        if text(row["Duplicate decision"]) == "Different specimen":
            record["human_duplicate_decision"] = "Different specimen"
        if prop not in PROPERTIES:
            record["context_measurements"].append({"property": prop, "value": json_value(row["Reported value"]), "unit": text(row["Reported unit"]), "candidate_id": candidate_id})
            decision["import_decision"] = "retained_context_not_plot_property"
            continue
        try:
            value, metadata = normalized_measurement(row, original)
        except ValueError as error:
            decision["import_decision"] = "held_ambiguous_numeric_expression"
            decision["reason"] = str(error)
            continue
        field = PROPERTIES[prop][0]
        release_property = {"youngs_modulus": "initial_modulus", "toughness": "work_of_rupture", "electrical_resistivity": "electrical_conductivity"}.get(prop, prop)
        if field in record["values"]:
            raise ValueError(f"Two approved values occupy {group_id}/{field}")
        record["values"][field] = value
        record["measurement_metadata"][release_property] = metadata
        decision["import_decision"] = "accepted"
        decision["record_id"] = record["record_id"]
        decision["canonical_value"] = value
        if metadata["raw_value_corrected"]:
            corrections.append({"candidate_id": candidate_id, "old_raw": text(original["raw_value"]),
                                "corrected_raw": text(row["Reported value"]), "canonical_value": value,
                                "workbook_normalized_value": json_value(row["Normalized value"])})
    records = [r for r in groups.values() if r["values"]]
    return {"schema_version": "cpt-curated-records-v1", "source_workbook": workbook.name,
            "source_sha256": hashlib.sha256(workbook.read_bytes()).hexdigest(),
            "source_candidate_sha256": hashlib.sha256(candidate_path.read_bytes()).hexdigest(),
            "status_counts": dict(counts), "import_counts": dict(Counter(d["import_decision"] for d in decisions)),
            "record_count": len(records), "corrected_values": corrections,
            "records": records, "decisions": decisions}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workbook", type=Path)
    parser.add_argument("--candidates", type=Path, default=CANDIDATES)
    parser.add_argument("--output", type=Path, default=ROOT / "data/curation/human_validated_records.json")
    args = parser.parse_args()
    payload = import_workbook(args.workbook, args.candidates)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(json.dumps(payload, indent=2, ensure_ascii=True) + "\n")
    print(json.dumps({k: payload[k] for k in ["status_counts", "import_counts", "record_count", "corrected_values"]}, indent=2))


if __name__ == "__main__":
    main()
