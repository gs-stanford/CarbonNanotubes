#!/usr/bin/env python3
"""Independent source-workbook and regression checks for the recent packets."""
import copy
import io
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import openpyxl
import pandas as pd

from build_combined_database import build_measurements, read_curated_records
from build_public_release import (build_public_measurements, build_public_records,
                                  duplicate_pair, release_decision)

ROOT = Path(__file__).resolve().parents[1]


def packet(name):
    return json.loads((ROOT / f"data/curation/primary_additions/{name}_2026.json").read_text())


class PrimaryAdmission(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        # Exercise the shared readers in isolation; never rebuild the release.
        with tempfile.TemporaryDirectory(prefix="sun-liu-admission-") as directory:
            fixture_root = Path(directory)
            additions = fixture_root / "data/curation/primary_additions"
            additions.mkdir(parents=True)
            for name in ("sun", "liu"):
                filename = f"{name}_2026.json"
                (additions / filename).symlink_to(ROOT / "data/curation/primary_additions" / filename)
            with patch("build_combined_database.ROOT", fixture_root):
                cls.loaded = read_curated_records()
        cls.measurements = build_measurements(copy.deepcopy(cls.loaded))
        enriched = []
        for original in cls.loaded:
            # Verified-DOI flags are projection-test fixtures, not new DOI checks.
            record = {**original, "publication_validation_status": "verified_crossref_openalex",
                      "doi_verified": original["doi_raw"], "publication_year_verified": 2026,
                      "publication_title_verified": "Projection fixture",
                      "publication_journal_verified": "Projection fixture"}
            record.update(release_decision(pd.Series(record)))
            enriched.append(record)
        cls.public_records = build_public_records(pd.DataFrame(enriched))
        cls.public_measurements = build_public_measurements(cls.measurements, cls.public_records)
        cls.csv_records = pd.read_csv(io.StringIO(cls.public_records.to_csv(index=False)))
        cls.csv_measurements = pd.read_csv(io.StringIO(cls.public_measurements.to_csv(index=False)))

    def test_sun_unique_controls_and_no_invented_transport(self):
        data = packet("sun")
        self.assertEqual(len(data["records"]), 19)
        self.assertEqual(sum(len(r["values"]) for r in data["records"]), 30)
        records = {r["sample_name"]: r for r in data["records"]}
        self.assertEqual(len(records), 19)
        self.assertNotIn("N-CNTFs", records)
        self.assertNotIn("MoCl5-CNTFs", records)
        self.assertNotIn("N/MoCl5-CNTFs-48 h", records)
        for r in records.values():
            self.assertNotIn("thermal_conductivity_W_mK", r["values"])
            self.assertNotIn("density_kg_m3", r["values"])
        self.assertEqual(records["MoCl5/N-CNTFs"]["values"], {"electrical_conductivity_S_m": 12632000})

    def test_sun_source_cells(self):
        book = openpyxl.load_workbook(ROOT / "outputs/fiber_discovery_20261022/sun_doping_audit/41467_2026_69498_MOESM3_ESM.xlsx", data_only=True)
        records = {r["sample_name"]: r for r in packet("sun")["records"]}
        for row in list(book["Supplementary Figure 3g"].values)[2:]:
            r = records[row[0]]
            self.assertAlmostEqual(r["values"]["tensile_strength_Pa"], row[1]*1e9)
            self.assertEqual(r["measurement_metadata"]["tensile_strength"]["uncertainty_value_reported"], row[2])
        for row in list(book["Supplementary Figure 4e"].values)[2:]:
            self.assertAlmostEqual(records[row[0]]["values"]["electrical_conductivity_S_m"], row[1]*1e6)
        book.close()

    def test_liu_source_values_and_scope(self):
        records = packet("liu")["records"]
        self.assertEqual(len(records), 10)
        self.assertEqual(sum(len(r["values"]) for r in records), 35)
        book = openpyxl.load_workbook(ROOT / "outputs/fiber_discovery_20261022/graphene_draw_audit/41563_2026_2733_MOESM10_ESM.xlsx", data_only=True)
        ext = openpyxl.load_workbook(ROOT / "outputs/fiber_discovery_20261022/graphene_draw_audit/41563_2026_2733_MOESM14_ESM.xlsx", data_only=True)
        for record, row in zip(records[:5], list(book["Fig.4f"].values)[1:]):
            self.assertEqual(record["values"]["tensile_strength_Pa"], row[3]*1e9)
            self.assertEqual(record["values"]["thermal_conductivity_W_mK"], row[1])
            self.assertEqual(record["measurement_metadata"]["density"]["extraction_method"], "vector_figure_digitization")
            self.assertNotIn("specific_strength_N_m_kg", record["values"])
        self.assertEqual({r["gauge_length_mm"] for r in records[5:]}, {5,10,15,20,25})
        for record, row in zip(records[5:], list(ext["E.D.Fig.6b"].values)[1:]):
            self.assertEqual(set(record["values"]), {"tensile_strength_Pa", "initial_modulus_Pa"})
            self.assertEqual(record["gauge_length_mm"], row[0])
            for prop, mean_column, sd_column in [("tensile_strength", 1, 2), ("initial_modulus", 3, 4)]:
                meta = record["measurement_metadata"][prop]
                self.assertEqual(meta["reported_value"], row[mean_column])
                self.assertEqual(meta["uncertainty_value_reported"], row[sd_column])
                self.assertEqual(meta["sample_size_n"], 3)
                self.assertEqual(meta["statistic_type"], "mean")
                self.assertEqual(meta["uncertainty_type"], "standard_deviation")
                self.assertEqual(meta["measurement_gauge_length_mm"], row[0])
        for record, row in zip(records[:5], list(ext["E.D.Fig.6d"].values)[1:]):
            meta = record["measurement_metadata"]["initial_modulus"]
            self.assertEqual(meta["reported_value"], row[1])
            self.assertEqual(meta["uncertainty_value_reported"], row[2])
        book.close()
        ext.close()

    def test_liu_five_mm_endpoint_is_a_separate_mechanical_set(self):
        records = {r["sample_name"]: r for r in packet("liu")["records"]}
        draw = records["Graphene fiber, draw ratio 10"]
        gauge = records["Graphene fiber, draw ratio 10, gauge 5 mm"]
        self.assertNotEqual(draw["record_id"], gauge["record_id"])
        self.assertNotEqual(draw["curated_state_id"], gauge["curated_state_id"])
        self.assertFalse(duplicate_pair(pd.Series(draw), pd.Series(gauge))[0])
        self.assertEqual(set(gauge["values"]), {"tensile_strength_Pa", "initial_modulus_Pa"})
        for prop, mean, draw_sd, gauge_sd in [("tensile_strength", 5.89, 0.14, 0.20),
                                             ("initial_modulus", 960, 55, 25)]:
            self.assertEqual(draw["measurement_metadata"][prop]["reported_value"], mean)
            self.assertEqual(gauge["measurement_metadata"][prop]["reported_value"], mean)
            self.assertEqual(draw["measurement_metadata"][prop]["uncertainty_value_reported"], draw_sd)
            self.assertEqual(gauge["measurement_metadata"][prop]["uncertainty_value_reported"], gauge_sd)
            for r in (draw, gauge):
                self.assertIn("specimen independence are unresolved",
                              r["measurement_metadata"][prop]["source_qualification"])
                self.assertEqual(r["specimen_linkage"], "single_source_row_unverified")

    def test_sun_dopant_methods_and_cui_source_correction(self):
        records = {r["sample_name"]: r for r in packet("sun")["records"]}
        for label, temperature in [("N/I2-CNTFs", 200), ("N/CuI-CNTFs", 750), ("N/FeCl3-CNTFs", 360)]:
            record = records[label]
            self.assertEqual(record["heat_treatment_C"], temperature)
            self.assertIsNone(record.get("condition_temperature_C"))
            for notes in (record["conditions_notes"],
                          record["measurement_metadata"]["electrical_conductivity"]["source_qualification"]):
                self.assertIn(f"annealed at {temperature} C", notes)
                self.assertIn("acetone-washed and dried at 60 C", notes)
                self.assertIn("not electrical measurement temperatures", notes)
                self.assertIn("does not isolate dopant chemistry alone", notes)
                self.assertNotIn("MoCl5/MoO3", notes)
        cui = records["N/CuI-CNTFs"]
        meta = cui["measurement_metadata"]["electrical_conductivity"]
        self.assertEqual(meta["reported_value"], 25.97)
        self.assertEqual(meta["uncertainty_value_reported"], 0.09)
        self.assertEqual(meta["uncertainty_value_canonical"], 90000)
        book = openpyxl.load_workbook(ROOT / "outputs/fiber_discovery_20261022/sun_doping_audit/41467_2026_69498_MOESM3_ESM.xlsx", data_only=True, read_only=True)
        self.assertEqual(book["Supplementary Figure 13a"]["A4"].value, "N/CuI-CNTFs")
        self.assertEqual(book["Supplementary Figure 13a"]["C4"].value, 0.09)
        book.close()
        for notes in (meta["source_qualification"], cui["conditions_notes"]):
            self.assertIn("25.97 +/- 0.09 MS/m", notes)
            self.assertIn("Supplementary Figure 13a!B4:C4", notes)
            self.assertIn("article.pdf p.5", notes)
            self.assertIn("was incorrect", notes)
            self.assertIn("Retain SD 0.09 MS/m", notes)

    def test_liu_conflicts_are_attached_to_affected_measurements(self):
        records = {r["sample_name"]: r for r in packet("liu")["records"]}
        expected = [
            ("Graphene fiber, draw ratio 1", "initial_modulus", ["38.1 +/- 67 GPa", "SD exceeds the mean"]),
            ("Graphene fiber, draw ratio 10", "initial_modulus", ["960 GPa", "963 GPa", "SI Table 3", "960 +/- 25 GPa", "960 +/- 55 GPa"]),
            ("Graphene fiber, draw ratio 10", "tensile_strength", ["5.89 +/- 0.14 GPa", "5.89 +/- 0.20 GPa"]),
            ("Graphene fiber, draw ratio 10", "electrical_conductivity", ["1.3301 MS/m", "1.35 MS/m", "SI Table 3"]),
        ]
        for label, prop, fragments in expected:
            record = records[label]
            for notes in (record["conditions_notes"], record["measurement_metadata"][prop]["source_qualification"]):
                for fragment in fragments:
                    self.assertIn(fragment, notes)
        for record in list(records.values())[:5]:
            thermal = record["measurement_metadata"]["thermal_conductivity"]["source_qualification"]
            self.assertIn("10^4 Pa", thermal)
            self.assertIn("inconsistency is unresolved", thermal)
            for prop in ("electrical_conductivity", "density"):
                meta = record["measurement_metadata"][prop]
                self.assertEqual(meta["extraction_method"], "vector_figure_digitization")
                self.assertIn("approximate figure coordinates", meta["source_qualification"])
                self.assertIn("not experimental SD", meta["source_qualification"])

    def test_qualifications_survive_loader_and_public_csv_projection(self):
        long = self.measurements.set_index(["record_id", "property"])
        public = self.csv_measurements.set_index(["record_id", "property"])
        records = self.csv_records.set_index("record_id")
        for name in ("sun", "liu"):
            for r in packet(name)["records"]:
                self.assertEqual(records.loc[r["record_id"], "conditions_notes"], r["conditions_notes"])
                for prop, meta in r["measurement_metadata"].items():
                    if "source_qualification" not in meta:
                        continue
                    key = (r["record_id"], prop)
                    self.assertEqual(long.loc[key, "source_qualification"], meta["source_qualification"])
                    self.assertEqual(public.loc[key, "source_qualification"], meta["source_qualification"])
                    self.assertNotIn("packet source_discrepancies", public.loc[key, "source_qualification"])
        self.assertEqual(len(self.csv_records), 29)
        self.assertEqual(len(self.csv_measurements[self.csv_measurements.source_file.eq("sun_2026.json")]), 30)
        liu = self.csv_measurements[self.csv_measurements.source_file.eq("liu_2026.json")]
        # The loader also computes five reciprocal-density values, not new observations.
        self.assertEqual(len(liu), 40)
        self.assertEqual(len(liu[~liu.property.eq("specific_volume")]), 35)
        self.assertFalse(self.csv_records.strict_comparison_ready.any())
        self.assertTrue(self.csv_records.primary_source_verification_status.eq(
            "ai_source_checked_not_human_adjudicated").all())
        self.assertFalse(liu.property.isin(["specific_strength", "specific_modulus",
            "specific_electrical_conductivity", "specific_thermal_conductivity"]).any())

    def test_per_property_n_is_not_overwritten(self):
        record = self.loaded[0].copy()
        record["sample_size_n"] = 20
        record["measurement_metadata_json"] = json.dumps({"density": {"sample_size_n": None},
            "tensile_strength": {"sample_size_n": 3}})
        record.update(density_kg_m3=1000, tensile_strength_Pa=1e9)
        rows = build_measurements([record]).set_index("property")
        self.assertTrue(pd.isna(rows.loc["density", "sample_size_n"]))
        self.assertEqual(rows.loc["tensile_strength", "sample_size_n"], 3)

    def test_equal_values_do_not_merge_explicitly_different_states(self):
        row = dict(material_family="CNT_or_CNT_hybrid", form_factor="fiber_yarn", doi_verified="10.1038/s41467-026-69498-7",
            tensile_strength_Pa=2e9, electrical_conductivity_S_m=1e7, sample_name="control")
        a, b = pd.Series({**row, "curated_state_id":"dose-1"}), pd.Series({**row, "curated_state_id":"dose-2"})
        self.assertFalse(duplicate_pair(a,b)[0])


if __name__ == "__main__":
    unittest.main()
